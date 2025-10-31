import google.generativeai as genai
from docx import Document
from fpdf import FPDF
import re
import PyPDF2
import io
import json

class CVParserAgent:
    def parse_cv(self, file_bytes):
        """Extracts text from a CV and uses Gemini to parse it into structured data."""
        text = self._extract_text(file_bytes)
        return self._parse_text_with_gemini(text)

    def _extract_text(self, file_bytes):
        """Extracts text from a PDF or DOCX file."""
        if file_bytes.name.endswith(".pdf"):
            return self._extract_text_from_pdf(file_bytes)
        elif file_bytes.name.endswith(".docx"):
            return self._extract_text_from_docx(file_bytes)
        else:
            return ""

    def _extract_text_from_pdf(self, file_bytes):
        """Extracts text from a PDF file."""
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes.read()))
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text

    def _extract_text_from_docx(self, file_bytes):
        """Extracts text from a DOCX file."""
        document = Document(io.BytesIO(file_bytes.read()))
        text = ""
        for para in document.paragraphs:
            text += para.text + "\n"
        return text

    def _parse_text_with_gemini(self, text):
        """Uses Gemini to parse the extracted text into a structured JSON object."""
        model = genai.GenerativeModel('gemini-1.5-flash')
        prompt = f"""
        Based on the following CV text, extract the user's information into a JSON object.
        The JSON object should have the following keys: "full_name", "contact_info", "linkedin_profile", "experience", "skills", "education".

        CV Text:
        {text}

        JSON Output:
        """
        response = model.generate_content(prompt)
        # Clean the response to ensure it's valid JSON
        cleaned_response = response.text.strip().replace("```json", "").replace("```", "")
        try:
            return json.loads(cleaned_response)
        except json.JSONDecodeError:
            return {}


class DataCollectionAgent:
    def collect_data(self, name, contact_info, linkedin_profile, experience, skills, education):
        """Gathers and structures the user's career data."""
        return {
            "name": name,
            "contact_info": contact_info,
            "linkedin_profile": linkedin_profile,
            "experience": experience,
            "skills": skills,
            "education": education
        }

class JDAnalysisAgent:
    def analyze(self, job_description):
        """Analyzes the job description to extract keywords and requirements."""
        # A more sophisticated approach would use NLP techniques like named entity recognition
        # and part-of-speech tagging to identify key skills and requirements.
        # For this implementation, we'll use a simple regex to extract potential keywords.
        keywords = re.findall(r'\b\w+\b', job_description.lower())
        return {
            "keywords": list(set(keywords)),
            "requirements": job_description
        }

class ResumeContentAgent:
    def generate(self, user_data, jd_analysis):
        """Generates tailored resume content based on the user's data and the job description analysis."""
        model = genai.GenerativeModel('gemini-1.5-flash')
        prompt = f"""
        Based on the following user data and job description, generate a professional resume.

        **User Data:**
        - Name: {user_data['name']}
        - Contact Info: {user_data['contact_info']}
        - LinkedIn: {user_data['linkedin_profile']}
        - Experience: {user_data['experience']}
        - Skills: {user_data['skills']}
        - Education: {user_data['education']}

        **Job Description:**
        {jd_analysis['requirements']}

        **Instructions:**
        - The resume should be tailored to the job description, highlighting the most relevant skills and experience.
        - Use a professional and clear format.
        - Emphasize achievements and quantifiable results.
        """
        response = model.generate_content(prompt)
        return response.text

class CoverLetterAgent:
    def generate(self, user_data, jd_analysis, job_title):
        """Generates a tailored cover letter."""
        model = genai.GenerativeModel('gemini-1.5-flash')
        prompt = f"""
        Based on the following user data and job description, generate a compelling cover letter.

        **User Data:**
        - Name: {user_data['name']}
        - Experience: {user_data['experience']}
        - Skills: {user_data['skills']}

        **Job Title:** {job_title}
        **Job Description:**
        {jd_analysis['requirements']}

        **Instructions:**
        - The cover letter should be addressed to the hiring manager (if not specified, use a generic greeting).
        - It should clearly articulate why the candidate is a good fit for the role, drawing direct connections between their experience and the job requirements.
        - Maintain a professional and enthusiastic tone.
        """
        response = model.generate_content(prompt)
        return response.text

class FormattingAgent:
    def to_pdf(self, text, filename):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 10, text)
        pdf.output(filename)

    def to_docx(self, text, filename):
        document = Document()
        document.add_paragraph(text)
        document.save(filename)
